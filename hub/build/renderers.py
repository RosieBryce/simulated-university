#!/usr/bin/env python3
"""
Renderers for the Stonegrove Hub build.

A source document is a Markdown file with YAML front-matter (see hub/README.md
for the schema). Each `format` has a renderer here that turns the body into a
real file: docx, pdf, pptx, xlsx, csv, txt, md, vtt, eml.

The Markdown dialect is deliberately small: headings, paragraphs, bullet and
numbered lists, pipe tables, bold and italic, horizontal rules, and a few
`::: directive` lines. `{{ key }}` placeholders are filled from values.yaml
before parsing; a key that is missing or null raises, because an in-world
document must never quote a figure the data does not have.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
from dataclasses import dataclass, field
from email.message import EmailMessage
from email.utils import format_datetime
from pathlib import Path
from typing import Any

import jinja2
import yaml

# ----------------------------------------------------------------------
# In-world dates
# ----------------------------------------------------------------------

YEAR_OFFSET = 973  # in-world 1053 -> real 2026, for file timestamps and email headers


def parse_inworld_date(value: Any) -> dt.date:
    """Accept 'YYYY-MM-DD' strings, or a date already parsed by YAML (which happens
    for valid-looking dates; in-world years are small but still valid)."""
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()
    y, m, d = (int(x) for x in s.split("-"))
    return dt.date(y, m, d)


def real_date(inworld: dt.date, hour: int = 10, minute: int = 30) -> dt.datetime:
    return dt.datetime(inworld.year + YEAR_OFFSET, inworld.month, inworld.day, hour, minute)


MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]


def inworld_long(d: dt.date) -> str:
    return f"{d.day} {MONTHS[d.month - 1]} {d.year}"


# ----------------------------------------------------------------------
# Values and cast
# ----------------------------------------------------------------------

class HubValues:
    """values.yaml wrapper. Null keys are removed so that quoting one raises."""

    def __init__(self, values: dict[str, Any]):
        self.values = {k: v for k, v in values.items() if v is not None and not k.startswith("_")}
        self.env = jinja2.Environment(undefined=jinja2.StrictUndefined, autoescape=False,
                                      keep_trailing_newline=True)
        self.env.filters["round"] = lambda x, n=1: round(float(x), n)
        self.env.filters["pp"] = lambda x: f"{float(x):.1f} percentage points"

    def fill(self, text: str, extra: dict[str, Any] | None = None) -> str:
        ctx = dict(self.values)
        if extra:
            ctx.update(extra)
        return self.env.from_string(text).render(**ctx)


class Cast:
    def __init__(self, cast_yaml: dict[str, Any]):
        self.people: dict[str, dict] = {}
        for p in cast_yaml.get("staff", []):
            self.people[p["id"]] = p
        for s in cast_yaml.get("students", []):
            self.people[s["id"]] = s
        self.institution = cast_yaml.get("institution", {})

    def name(self, cast_id: str | None) -> str:
        if not cast_id:
            return ""
        p = self.people.get(cast_id)
        if not p:
            raise KeyError(f"unknown cast id {cast_id!r}")
        return p["name"]

    def short(self, cast_id: str) -> str:
        """Name without honorific, for 'Modified By' columns and email display."""
        n = self.name(cast_id)
        return re.sub(r"^(Professor|Dr|Sister)\s+", "", n)

    def email(self, cast_id: str) -> str:
        p = self.people[cast_id]
        return p.get("email") or f"{cast_id.replace('-', '.')}@{self.institution.get('email_domain', 'stonegrove.ac.example')}"

    def title(self, cast_id: str) -> str:
        return self.people[cast_id].get("title", "")


# ----------------------------------------------------------------------
# Source files
# ----------------------------------------------------------------------

FRONT = re.compile(r"\A---\s*\n(.*?)\n---\s*\n?", re.S)


@dataclass
class SourceDoc:
    path: Path
    meta: dict[str, Any]
    body: str

    @property
    def fmt(self) -> str:
        return str(self.meta.get("format", "docx")).lower()

    @property
    def stem(self) -> str:
        return self.path.stem


def read_source(path: Path) -> SourceDoc:
    text = path.read_text(encoding="utf-8")
    m = FRONT.match(text)
    if not m:
        raise ValueError(f"{path}: no front-matter")
    meta = yaml.safe_load(m.group(1)) or {}
    body = text[m.end():]
    return SourceDoc(path=path, meta=meta, body=body)


# ----------------------------------------------------------------------
# Mini Markdown parser
# ----------------------------------------------------------------------

@dataclass
class Block:
    kind: str                 # heading | para | bullets | numbers | table | rule | directive | code
    text: str = ""
    level: int = 0
    items: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    args: list[str] = field(default_factory=list)
    lang: str = ""


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")


def inline_runs(text: str) -> list[tuple[str, bool, bool, bool]]:
    """Split a line into (text, bold, italic, code) runs."""
    runs = []
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            runs.append((part[2:-2], True, False, False))
        elif part.startswith("`") and part.endswith("`"):
            runs.append((part[1:-1], False, False, True))
        elif part.startswith("*") and part.endswith("*"):
            runs.append((part[1:-1], False, True, False))
        else:
            runs.append((part, False, False, False))
    return runs


def plain(text: str) -> str:
    return "".join(r[0] for r in inline_runs(text))


def parse_markdown(body: str) -> list[Block]:
    lines = body.splitlines()
    blocks: list[Block] = []
    i = 0
    para: list[str] = []

    def flush():
        nonlocal para
        if para:
            blocks.append(Block("para", " ".join(s.strip() for s in para)))
            para = []

    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s:
            flush(); i += 1; continue
        if s.startswith("```"):
            flush()
            lang = s[3:].strip()
            j = i + 1
            code: list[str] = []
            while j < len(lines) and not lines[j].strip().startswith("```"):
                code.append(lines[j]); j += 1
            blocks.append(Block("code", "\n".join(code), lang=lang))
            i = j + 1; continue
        if s.startswith(":::"):
            flush()
            blocks.append(Block("directive", args=s[3:].split()))
            i += 1; continue
        if re.match(r"^#{1,6}\s", s):
            flush()
            level = len(s) - len(s.lstrip("#"))
            blocks.append(Block("heading", s[level:].strip(), level=level))
            i += 1; continue
        if re.match(r"^(-{3,}|\*{3,})$", s):
            flush(); blocks.append(Block("rule")); i += 1; continue
        if re.match(r"^[-*]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i])); i += 1
            blocks.append(Block("bullets", items=items)); continue
        if re.match(r"^\d+[.)]\s+", s):
            flush()
            items = []
            while i < len(lines) and re.match(r"^\s*\d+[.)]\s+", lines[i]):
                items.append(re.sub(r"^\s*\d+[.)]\s+", "", lines[i])); i += 1
            blocks.append(Block("numbers", items=items)); continue
        if s.startswith("|"):
            flush()
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(Block("table", rows=rows)); continue
        para.append(line); i += 1
    flush()
    return blocks


# ----------------------------------------------------------------------
# Rendering context
# ----------------------------------------------------------------------

@dataclass
class RenderContext:
    cast: Cast
    values: HubValues
    doc: SourceDoc
    out_path: Path
    label: str | None
    built_files: dict[str, Path] = field(default_factory=dict)   # source rel path -> built file, for eml attachments

    @property
    def author_name(self) -> str:
        return self.cast.short(self.doc.meta["author"]) if self.doc.meta.get("author") else ""

    @property
    def modified_by_name(self) -> str:
        mb = self.doc.meta.get("modified_by") or self.doc.meta.get("author")
        return self.cast.short(mb) if mb else ""

    @property
    def created(self) -> dt.date:
        return parse_inworld_date(self.doc.meta.get("created") or self.doc.meta.get("modified"))

    @property
    def modified(self) -> dt.date:
        return parse_inworld_date(self.doc.meta.get("modified") or self.doc.meta.get("created"))

    @property
    def title(self) -> str:
        return self.doc.meta.get("title") or self.doc.stem


def filled_body(ctx: RenderContext) -> str:
    extra = {
        "author_name": ctx.author_name,
        "today": inworld_long(ctx.modified),
        "institution": ctx.cast.institution.get("name", "Stonegrove University"),
    }
    return ctx.values.fill(ctx.doc.body, extra)


# ----------------------------------------------------------------------
# DOCX
# ----------------------------------------------------------------------

def render_docx(ctx: RenderContext) -> Path:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    body = filled_body(ctx)
    blocks = parse_markdown(body)
    d = docx.Document()

    style = d.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)
    letterhead = "current"
    footer_mode = None
    for b in blocks:
        if b.kind == "directive" and b.args and b.args[0] == "letterhead" and len(b.args) > 1:
            letterhead = b.args[1]
        if b.kind == "directive" and b.args and b.args[0] == "footer" and len(b.args) > 1:
            footer_mode = b.args[1]

    inst = ctx.cast.institution
    section = d.sections[0]
    header = section.header.paragraphs[0]
    if letterhead == "old":
        header.text = inst.get("former_name", "Stonegrove College of the Vale")
        header.runs[0].font.name = "Georgia"
        header.runs[0].font.color.rgb = RGBColor(0x7A, 0x1F, 0x2E)
    elif letterhead == "none":
        header.text = ""
    else:
        header.text = inst.get("name", "Stonegrove University")
        header.runs[0].font.name = "Georgia"
        header.runs[0].font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    footer_bits = []
    if ctx.label:
        footer_bits.append(ctx.label.upper())
    if letterhead == "old" or footer_mode == "old":
        footer_bits.append(inst.get("former_name", ""))
    elif footer_mode == "current":
        footer_bits.append(inst.get("name", ""))
    footer.text = "    ".join(footer_bits)
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_runs(par, text):
        for t, bold, italic, code in inline_runs(text):
            r = par.add_run(t)
            r.bold = bold or None
            r.italic = italic or None
            if code:
                r.font.name = "Consolas"

    for b in blocks:
        if b.kind == "heading":
            d.add_heading(plain(b.text), level=min(b.level, 4))
        elif b.kind == "para":
            add_runs(d.add_paragraph(), b.text)
        elif b.kind == "bullets":
            for it in b.items:
                add_runs(d.add_paragraph(style="List Bullet"), it)
        elif b.kind == "numbers":
            for it in b.items:
                add_runs(d.add_paragraph(style="List Number"), it)
        elif b.kind == "table" and b.rows:
            ncol = max(len(r) for r in b.rows)
            t = d.add_table(rows=0, cols=ncol)
            t.style = "Table Grid"
            for ri, row in enumerate(b.rows):
                cells = t.add_row().cells
                for ci in range(ncol):
                    txt = row[ci] if ci < len(row) else ""
                    p = cells[ci].paragraphs[0]
                    add_runs(p, txt)
                    if ri == 0:
                        for r in p.runs:
                            r.bold = True
        elif b.kind == "rule":
            d.add_paragraph("_" * 40)
        elif b.kind == "code":
            p = d.add_paragraph()
            r = p.add_run(b.text)
            r.font.name = "Consolas"
            r.font.size = Pt(9)

    cp = d.core_properties
    cp.title = ctx.title
    cp.author = ctx.author_name
    cp.last_modified_by = ctx.modified_by_name
    cp.created = real_date(ctx.created)
    cp.modified = real_date(ctx.modified)
    if ctx.label:
        cp.category = ctx.label
    d.save(ctx.out_path)
    return ctx.out_path


# ----------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------

def render_pdf(ctx: RenderContext) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    body = filled_body(ctx)
    blocks = parse_markdown(body)
    ss = getSampleStyleSheet()
    styles = {
        "body": ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=10.5, leading=14),
        1: ParagraphStyle("h1", parent=ss["Heading1"], fontName="Times-Bold", fontSize=18, spaceAfter=8),
        2: ParagraphStyle("h2", parent=ss["Heading2"], fontName="Times-Bold", fontSize=14, spaceAfter=6),
        3: ParagraphStyle("h3", parent=ss["Heading3"], fontName="Helvetica-Bold", fontSize=11.5),
    }

    def rl(text: str) -> str:
        out = ""
        for t, bold, italic, code in inline_runs(text):
            t = t.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if bold:
                t = f"<b>{t}</b>"
            if italic:
                t = f"<i>{t}</i>"
            if code:
                t = f"<font face='Courier'>{t}</font>"
            out += t
        return out

    inst = ctx.cast.institution
    letterhead = "current"
    for b in blocks:
        if b.kind == "directive" and b.args and b.args[0] == "letterhead" and len(b.args) > 1:
            letterhead = b.args[1]
    head_name = inst.get("former_name") if letterhead == "old" else inst.get("name", "Stonegrove University")

    def on_page(canvas, doc_):
        canvas.saveState()
        if letterhead != "none":
            canvas.setFont("Times-Roman", 9)
            canvas.setFillColor(colors.HexColor("#7A1F2E" if letterhead == "old" else "#1A1A2E"))
            canvas.drawRightString(A4[0] - 20 * mm, A4[1] - 15 * mm, head_name)
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        bits = [ctx.label.upper()] if ctx.label else []
        bits.append(f"Page {doc_.page}")
        canvas.drawCentredString(A4[0] / 2, 12 * mm, "    ".join(bits))
        canvas.restoreState()

    story: list = []
    for b in blocks:
        if b.kind == "heading":
            story.append(Paragraph(rl(b.text), styles.get(b.level, styles[3])))
        elif b.kind == "para":
            story.append(Paragraph(rl(b.text), styles["body"]))
            story.append(Spacer(1, 4))
        elif b.kind in ("bullets", "numbers"):
            story.append(ListFlowable([ListItem(Paragraph(rl(i), styles["body"])) for i in b.items],
                                      bulletType="bullet" if b.kind == "bullets" else "1", leftIndent=14))
            story.append(Spacer(1, 4))
        elif b.kind == "table" and b.rows:
            data = [[Paragraph(rl(c), styles["body"]) for c in row] for row in b.rows]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDE8DC")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))
        elif b.kind == "rule":
            story.append(Spacer(1, 10))
        elif b.kind == "directive" and b.args and b.args[0] == "pagebreak":
            story.append(PageBreak())
        elif b.kind == "code":
            story.append(Paragraph(rl(b.text).replace("\n", "<br/>"), ParagraphStyle("code", parent=styles["body"], fontName="Courier", fontSize=8.5)))

    doc_ = SimpleDocTemplate(str(ctx.out_path), pagesize=A4, title=ctx.title, author=ctx.author_name,
                             leftMargin=22 * mm, rightMargin=22 * mm, topMargin=25 * mm, bottomMargin=22 * mm)
    doc_.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return ctx.out_path


# ----------------------------------------------------------------------
# PPTX
# ----------------------------------------------------------------------

def render_pptx(ctx: RenderContext) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    body = filled_body(ctx)
    theme = "stonegrove"
    slides_md = [s for s in re.split(r"\n-{3,}\n", body)]
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    palette = {
        "stonegrove": dict(bg=RGBColor(0xF5, 0xF0, 0xE0), title=RGBColor(0x1A, 0x1A, 0x2E), text=RGBColor(0x2E, 0x2E, 0x42),
                           accent=RGBColor(0xC9, 0xA8, 0x4C), tfont="Georgia", bfont="Arial"),
        "jaskel": dict(bg=RGBColor(0xFF, 0xF4, 0xB3), title=RGBColor(0xFF, 0x3C, 0xA6), text=RGBColor(0x00, 0x66, 0xCC),
                       accent=RGBColor(0x33, 0xCC, 0x33), tfont="Comic Sans MS", bfont="Impact"),
        "old": dict(bg=RGBColor(0xFF, 0xFF, 0xFF), title=RGBColor(0x7A, 0x1F, 0x2E), text=RGBColor(0x33, 0x33, 0x33),
                    accent=RGBColor(0x7A, 0x1F, 0x2E), tfont="Times New Roman", bfont="Times New Roman"),
    }
    jaskel_fonts = ["Comic Sans MS", "Impact", "Papyrus", "Curlz MT", "Jokerman", "Segoe Script"]

    for si, smd in enumerate(slides_md):
        blocks = parse_markdown(smd)
        for b in blocks:
            if b.kind == "directive" and b.args and b.args[0] == "theme" and len(b.args) > 1:
                theme = b.args[1]
        pal = palette.get(theme, palette["stonegrove"])
        slide = prs.slides.add_slide(blank)
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = pal["bg"]
        title_text = next((plain(b.text) for b in blocks if b.kind == "heading"), ctx.title if si == 0 else "")
        tb = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(12), Inches(1.2))
        p = tb.text_frame.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(36 if si else 44)
        p.font.bold = True
        p.font.name = pal["tfont"] if theme != "jaskel" else jaskel_fonts[si % len(jaskel_fonts)]
        p.font.color.rgb = pal["title"]
        body_box = slide.shapes.add_textbox(Inches(0.9), Inches(1.9), Inches(11.5), Inches(5))
        tf = body_box.text_frame
        tf.word_wrap = True
        first = True
        for b in blocks:
            if b.kind == "heading" and plain(b.text) == title_text:
                continue
            lines: list[tuple[str, int]] = []
            if b.kind == "para":
                lines.append((plain(b.text), 0))
            elif b.kind in ("bullets", "numbers"):
                lines.extend((("• " if b.kind == "bullets" else f"{n}. ") + plain(i), 1) for n, i in enumerate(b.items, 1))
            elif b.kind == "heading":
                lines.append((plain(b.text), 0))
            elif b.kind == "table":
                lines.extend(("   ".join(r), 1) for r in b.rows)
            for text, lvl in lines:
                para = tf.paragraphs[0] if first else tf.add_paragraph()
                first = False
                para.text = text
                para.level = lvl
                para.font.size = Pt(20 if lvl == 0 else 18)
                para.font.name = pal["bfont"] if theme != "jaskel" else jaskel_fonts[(si + lvl + len(text)) % len(jaskel_fonts)]
                para.font.color.rgb = pal["text"]
        if theme == "jaskel":
            shp = slide.shapes.add_shape(1, Inches(11.3), Inches(5.9), Inches(1.6), Inches(1.2))
            shp.fill.solid()
            shp.fill.fore_color.rgb = pal["accent"]
            shp.text_frame.text = "pencil tree"
            shp.shadow.inherit = True

    cp = prs.core_properties
    cp.title = ctx.title
    cp.author = ctx.author_name
    cp.last_modified_by = ctx.modified_by_name
    cp.created = real_date(ctx.created)
    cp.modified = real_date(ctx.modified)
    prs.save(ctx.out_path)
    return ctx.out_path


# ----------------------------------------------------------------------
# XLSX / CSV
# ----------------------------------------------------------------------

def _sheets_from_body(body: str) -> list[tuple[str, list[list[str]]]]:
    """Fenced blocks tagged `sheet:Name` (or untagged) hold CSV."""
    sheets = []
    for m in re.finditer(r"```(?:sheet:)?([^\n`]*)\n(.*?)```", body, re.S):
        name = (m.group(1).strip() or "Sheet1")[:31]
        rows = list(csv.reader(io.StringIO(m.group(2).strip("\n"))))
        sheets.append((name, rows))
    return sheets


def _coerce(cell: str):
    if cell == "":
        return None
    if cell.startswith("="):
        return cell
    try:
        if re.fullmatch(r"-?\d+", cell):
            return int(cell)
        if re.fullmatch(r"-?\d+\.\d+", cell):
            return float(cell)
    except ValueError:
        pass
    return cell


def render_xlsx(ctx: RenderContext) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    body = filled_body(ctx)
    sheets = _sheets_from_body(body)
    if not sheets:
        raise ValueError(f"{ctx.doc.path}: xlsx source has no fenced CSV blocks")
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets:
        ws = wb.create_sheet(title=name)
        for r in rows:
            ws.append([_coerce(c) for c in r])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="EDE8DC")
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
    wb.properties.title = ctx.title
    wb.properties.creator = ctx.author_name
    wb.properties.lastModifiedBy = ctx.modified_by_name
    wb.properties.created = real_date(ctx.created)
    wb.properties.modified = real_date(ctx.modified)
    wb.save(ctx.out_path)
    return ctx.out_path


def render_csv(ctx: RenderContext) -> Path:
    body = filled_body(ctx)
    sheets = _sheets_from_body(body)
    if not sheets:
        raise ValueError(f"{ctx.doc.path}: csv source has no fenced CSV block")
    with open(ctx.out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        for r in sheets[0][1]:
            w.writerow(r)
    return ctx.out_path


# ----------------------------------------------------------------------
# TXT / MD
# ----------------------------------------------------------------------

def render_txt(ctx: RenderContext) -> Path:
    body = filled_body(ctx)
    if ctx.doc.fmt == "txt":
        # strip the lightest markdown so it reads as plain text
        body = re.sub(r"^#{1,6}\s*", "", body, flags=re.M)
        body = body.replace("**", "")
    ctx.out_path.write_text(body, encoding="utf-8", newline="\r\n")
    return ctx.out_path


# ----------------------------------------------------------------------
# VTT transcripts
# ----------------------------------------------------------------------

CUE = re.compile(r"^\[(\d{2}):(\d{2}):(\d{2})\]\s+([a-z0-9-]+|[A-Z][\w ]*?):\s*(.*)$")


def _ts(h: int, m: int, s: int, ms: int = 0) -> str:
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def parse_transcript(body: str, cast: Cast, mangle: dict[str, str] | None) -> list[tuple[int, str, str]]:
    """Return (seconds, speaker display, text). Speaker may be a cast id or a literal label."""
    out = []
    for line in body.splitlines():
        line = line.rstrip()
        if not line or line.startswith("#") or line.startswith(":::"):
            continue
        m = CUE.match(line)
        if not m:
            # continuation of previous cue
            if out:
                sec, sp, txt = out[-1]
                out[-1] = (sec, sp, txt + " " + line.strip())
            continue
        h, mi, s, who, text = m.groups()
        secs = int(h) * 3600 + int(mi) * 60 + int(s)
        if who in cast.people:
            disp = (mangle or {}).get(who) or cast.short(who)
        else:
            disp = who
        out.append((secs, disp, text))
    return out


def render_vtt(ctx: RenderContext) -> Path:
    body = filled_body(ctx)
    mangle = ctx.doc.meta.get("mangle") or {}
    cues = parse_transcript(body, ctx.cast, mangle)
    lines = ["WEBVTT", f"NOTE {ctx.title}", f"NOTE Recorded {inworld_long(ctx.modified)}. Auto-generated transcript.", ""]
    for i, (sec, who, text) in enumerate(cues):
        end = cues[i + 1][0] if i + 1 < len(cues) else sec + 6
        if end <= sec:
            end = sec + 3
        lines.append(f"{i + 1}")
        lines.append(f"{_ts(sec // 3600, sec % 3600 // 60, sec % 60)} --> {_ts(end // 3600, end % 3600 // 60, end % 60)}")
        lines.append(f"<v {who}>{text}")
        lines.append("")
    ctx.out_path.write_text("\n".join(lines), encoding="utf-8", newline="\n")
    return ctx.out_path


def render_transcript_docx(ctx: RenderContext) -> Path:
    """A docx of the same transcript, as Teams exports it: time, speaker, text."""
    import docx
    from docx.shared import Pt

    body = filled_body(ctx)
    mangle = ctx.doc.meta.get("mangle") or {}
    cues = parse_transcript(body, ctx.cast, mangle)
    d = docx.Document()
    d.styles["Normal"].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(10.5)
    d.add_heading(ctx.title, level=1)
    d.add_paragraph(f"Transcript, {inworld_long(ctx.modified)}. Generated automatically from the meeting recording. "
                    "Speaker names and words may be inaccurate.")
    for sec, who, text in cues:
        p = d.add_paragraph()
        r = p.add_run(f"{sec // 3600:02d}:{sec % 3600 // 60:02d}:{sec % 60:02d}  ")
        r.font.size = Pt(8.5)
        r2 = p.add_run(f"{who}: ")
        r2.bold = True
        p.add_run(text)
    cp = d.core_properties
    cp.title = ctx.title
    cp.author = "Meeting recording"
    cp.created = real_date(ctx.created)
    cp.modified = real_date(ctx.modified)
    d.save(ctx.out_path)
    return ctx.out_path


# ----------------------------------------------------------------------
# EML
# ----------------------------------------------------------------------

HDR = re.compile(r"^(From|To|Cc|Date|Subject|Attachments):\s*(.*)$", re.I)


def render_eml(ctx: RenderContext) -> Path:
    body = filled_body(ctx)
    headers: dict[str, str] = {}
    lines = body.splitlines()
    i = 0
    while i < len(lines) and lines[i].strip():
        m = HDR.match(lines[i])
        if not m:
            break
        headers[m.group(1).lower()] = m.group(2).strip()
        i += 1
    text = "\n".join(lines[i:]).strip("\n") + "\n"

    def addr(spec: str) -> str:
        """Cast ids become 'Name <email>'; anything else passes through."""
        parts = []
        for a in [x.strip() for x in spec.split(",") if x.strip()]:
            if a in ctx.cast.people:
                parts.append(f"{ctx.cast.short(a)} <{ctx.cast.email(a)}>")
            else:
                parts.append(a)
        return ", ".join(parts)

    msg = EmailMessage()
    msg["From"] = addr(headers.get("from", ctx.doc.meta.get("author", "")))
    msg["To"] = addr(headers.get("to", ""))
    if headers.get("cc"):
        msg["Cc"] = addr(headers["cc"])
    msg["Subject"] = headers.get("subject", ctx.title)
    when = parse_inworld_date(headers.get("date") or ctx.doc.meta.get("modified"))
    msg["Date"] = format_datetime(real_date(when, 9, 41))
    msg["X-Stonegrove-InWorld-Date"] = inworld_long(when)
    msg.set_content(text)
    for att in [a.strip() for a in headers.get("attachments", "").split(",") if a.strip()]:
        built = ctx.built_files.get(att)
        if not built:
            raise FileNotFoundError(f"{ctx.doc.path}: attachment {att!r} not built before this email")
        data = Path(built).read_bytes()
        ext = Path(built).suffix.lower().lstrip(".")
        maintype, subtype = {
            "docx": ("application", "vnd.openxmlformats-officedocument.wordprocessingml.document"),
            "xlsx": ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            "pptx": ("application", "vnd.openxmlformats-officedocument.presentationml.presentation"),
            "pdf": ("application", "pdf"), "png": ("image", "png"), "jpg": ("image", "jpeg"),
        }.get(ext, ("application", "octet-stream"))
        msg.add_attachment(data, maintype=maintype, subtype=subtype, filename=Path(built).name)
    ctx.out_path.write_bytes(msg.as_bytes())
    return ctx.out_path


# ----------------------------------------------------------------------
RENDERERS = {
    "docx": render_docx,
    "pdf": render_pdf,
    "pptx": render_pptx,
    "xlsx": render_xlsx,
    "csv": render_csv,
    "txt": render_txt,
    "md": render_txt,
    "vtt": render_vtt,
    "eml": render_eml,
}


def render(ctx: RenderContext) -> Path:
    fmt = ctx.doc.fmt
    if fmt not in RENDERERS:
        raise ValueError(f"{ctx.doc.path}: no renderer for format {fmt!r}")
    return RENDERERS[fmt](ctx)
